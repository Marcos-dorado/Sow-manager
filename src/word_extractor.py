"""
Módulo para extraer datos de documentos Word (.docx) de SOWs CHUBB-NEORIS.
Extrae: fecha inicio, fecha fin, valor COP y análisis completo del SOW.
"""

import io
import re
from datetime import datetime

MESES_ES = {
    "enero": 1, "febrero": 2, "marzo": 3, "abril": 4,
    "mayo": 5, "junio": 6, "julio": 7, "agosto": 8,
    "septiembre": 9, "octubre": 10, "noviembre": 11, "diciembre": 12,
}

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def parsear_fecha(texto):
    if not texto or not texto.strip():
        return None
    texto = texto.strip()
    for fmt in ["%m/%d/%Y", "%d/%m/%Y", "%m-%d-%Y",
                "%d-%m-%Y", "%Y-%m-%d", "%d/%m/%y", "%m/%d/%y"]:
        try:
            return datetime.strptime(texto, fmt)
        except ValueError:
            continue
    match = re.search(
        r"(\d{1,2})\s+de\s+(\w+)\s+de\s+(\d{4})", texto, re.IGNORECASE)
    if match:
        dia  = int(match.group(1))
        mes  = MESES_ES.get(match.group(2).lower())
        anio = int(match.group(3))
        if mes:
            try:
                return datetime(anio, mes, dia)
            except ValueError:
                pass
    return None


def limpiar_monto(texto):
    if not texto:
        return None
    texto = re.sub(r"COL\$|\$|\s", "", texto, flags=re.IGNORECASE)
    texto = texto.replace(",", "")
    try:
        valor = float(texto)
        return valor if valor > 0 else None
    except ValueError:
        return None


def extraer_sdt_por_tag(root, tag_name):
    for sdt in root.iter(f"{{{W}}}sdt"):
        sdtPr = sdt.find(f"{{{W}}}sdtPr")
        if sdtPr is None:
            continue
        tag_el = sdtPr.find(f"{{{W}}}tag")
        if tag_el is None:
            continue
        val = tag_el.get(f"{{{W}}}val", "")
        if val.lower() != tag_name.lower():
            continue
        sdtContent = sdt.find(f"{{{W}}}sdtContent")
        if sdtContent is None:
            continue
        textos = []
        for t in sdtContent.iter(f"{{{W}}}t"):
            if t.text:
                textos.append(t.text)
        texto = "".join(textos).strip()
        if texto:
            return texto
    return None


def extraer_texto_parrafos(doc):
    return "\n".join(p.text for p in doc.paragraphs)


def extraer_datos_de_tablas(tablas, numero_sow=None):
    resultado = {
        "fecha_inicio": None,
        "fecha_fin":    None,
        "valor_cop":    None,
        "fuente":       "",
    }
    numero_sow_upper = (numero_sow or "").strip().upper()

    for tabla in tablas:
        for i, fila in enumerate(tabla):
            fila_texto = " | ".join(str(c) for c in fila)
            fila_lower = fila_texto.lower()

            if "dt/sc" in fila_lower and "fecha inicio" in fila_lower:
                for fila_data in tabla[i + 1:]:
                    if len(fila_data) < 3:
                        continue
                    sow_celda = fila_data[0].strip().upper()
                    if numero_sow_upper and sow_celda == numero_sow_upper:
                        fi = parsear_fecha(fila_data[1])
                        ff = parsear_fecha(fila_data[2])
                        if fi and not resultado["fecha_inicio"]:
                            resultado["fecha_inicio"] = fi
                        if ff and not resultado["fecha_fin"]:
                            resultado["fecha_fin"] = ff
                        if len(fila_data) >= 5 and not resultado["valor_cop"]:
                            valor = limpiar_monto(fila_data[4])
                            if valor:
                                resultado["valor_cop"] = valor
                        break

            if ("costo tope total" in fila_lower or
                    "no podrá exceder" in fila_lower):
                for celda in fila:
                    if "COL$" in celda or "$" in celda:
                        valor = limpiar_monto(celda)
                        if valor and not resultado["valor_cop"]:
                            resultado["valor_cop"] = valor

            if not resultado["valor_cop"]:
                if (fila_texto.strip().startswith("Total (1)") or
                        fila_texto.strip().startswith("Total(1)")):
                    for celda in fila[1:]:
                        if "COL$" in celda or "$" in celda:
                            valor = limpiar_monto(celda)
                            if valor:
                                resultado["valor_cop"] = valor
                                break

    return resultado


# ══════════════════════════════════════════════════════════════
#  ANÁLISIS COMPLETO — 6 puntos
# ══════════════════════════════════════════════════════════════

def _validar_contrato_marco(texto, doc):
    """
    Punto 1: El contrato marco puede estar en content controls (SDT)
    o en el texto de párrafos. Como los campos están vacíos en el doc,
    buscamos en los SDT del XML directamente.
    """
    root = doc.element.body

    # Recoger TODO el texto incluyendo SDTs
    textos_sdt = []
    for sdt in root.iter(f"{{{W}}}sdt"):
        sdtContent = sdt.find(f"{{{W}}}sdtContent")
        if sdtContent is not None:
            for t in sdtContent.iter(f"{{{W}}}t"):
                if t.text and t.text.strip():
                    textos_sdt.append(t.text.strip())

    texto_completo = texto + " " + " ".join(textos_sdt)

    # Buscar fecha 14 de abril de 2020
    tiene_fecha = bool(re.search(
        r"14\s+de\s+abril\s+de\s+2020",
        texto_completo, re.IGNORECASE))

    # Buscar Chubb Seguros Mexico
    tiene_chubb = bool(re.search(
        r"chubb\s+seguros\s+m[eé]xico",
        texto_completo, re.IGNORECASE))

    # Buscar NEORIS Colombia
    tiene_neoris = bool(re.search(
        r"neoris\s+colombia",
        texto_completo, re.IGNORECASE))

    # Si no encontramos en texto visible, buscar número de acuerdo
    # Acuerdo 2020-NME-14042020-1 siempre está referenciado
    tiene_acuerdo = bool(re.search(
        r"14042020|14\s*de\s*abril\s*de\s*2020|2020-NME",
        texto_completo, re.IGNORECASE))

    # El contrato marco se valida si encontramos el acuerdo o la fecha
    ok = (tiene_fecha or tiene_acuerdo) and tiene_neoris

    detalle = []
    if tiene_fecha or tiene_acuerdo:
        detalle.append("✅ Referencia al contrato del 14 de abril de 2020")
    else:
        detalle.append("❌ No se encontró referencia al contrato marco")

    if tiene_chubb:
        detalle.append("✅ Chubb Seguros México referenciado")
    else:
        # En estos docs Chubb está como campo vacío — lo marcamos como info
        detalle.append("⚠️ Chubb Seguros México (campo en blanco en el doc)")

    if tiene_neoris:
        detalle.append("✅ NEORIS Colombia referenciado")
    else:
        detalle.append("❌ No se encontró referencia a NEORIS Colombia")

    return {"ok": ok, "detalle": detalle}


def _validar_tipo_servicio(texto):
    """
    Punto 2: Verifica que sea prestación de servicios.
    Texto real: 'El Prestador proveerá los servicios de asignación'
    """
    patron_ok = re.search(
        r"prover[aá]\s+los\s+servicios|"
        r"prestaci[oó]n\s+de\s+servicios|"
        r"servicios\s+de\s+asignaci[oó]n|"
        r"asignaci[oó]n\s+de\s+uno\s+o\s+varios\s+consultores",
        texto, re.IGNORECASE)

    patron_malo = re.search(
        r"\bventa\s+de\b|\blicencia\s+de\s+software\b|"
        r"\bproducto\s+de\s+software\b",
        texto, re.IGNORECASE)

    ok = bool(patron_ok) and not bool(patron_malo)

    if ok:
        detalle = "✅ Confirma prestación de servicios de consultoría"
    elif patron_malo:
        detalle = (f"❌ Se detectó posible venta/licencia: "
                   f"'{patron_malo.group()}'")
    else:
        detalle = "❌ No se encontró la cláusula de prestación de servicios"

    return {"ok": ok, "detalle": detalle}


def _extraer_cargos(doc):
    """
    Punto 3: Extrae cargos desde párrafos Y desde tablas.

    Formatos reales encontrados:
    - En párrafo: '1 (UN) Arquitecto de Soluciones Senior'
    - En tabla:   '1 (UN) QA2-Automatización QA\n"descripción..."'
                  '1 (UN) D6-Desarrollo Angular\n"descripción..."'
    """
    cargos  = []
    vistos  = set()

    patron = re.compile(
        r"(\d+)\s*\((?:UN|DOS|TRES|CUATRO|CINCO|UNA)\)\s+"
        r"([A-Z0-9][^\n\"]{2,80}?)(?=\n|\"|$)",
        re.IGNORECASE | re.MULTILINE
    )

    # Buscar en párrafos
    for p in doc.paragraphs:
        texto = p.text.strip()
        if not texto:
            continue
        for m in patron.finditer(texto):
            cantidad = m.group(1)
            cargo    = m.group(2).strip().rstrip("-").strip()
            key      = cargo.upper()[:40]
            if key not in vistos and len(cargo) > 3:
                vistos.add(key)
                cargos.append({
                    "cantidad": cantidad,
                    "cargo":    cargo,
                })

    # Buscar en tablas (cada celda)
    for tabla in doc.tables:
        for fila in tabla.rows:
            for celda in fila.cells:
                texto = celda.text.strip()
                if not texto:
                    continue
                for m in patron.finditer(texto):
                    cantidad = m.group(1)
                    cargo    = m.group(2).strip().rstrip("-").strip()
                    key      = cargo.upper()[:40]
                    if key not in vistos and len(cargo) > 3:
                        vistos.add(key)
                        cargos.append({
                            "cantidad": cantidad,
                            "cargo":    cargo,
                        })

    return cargos


def _validar_vigencia(fecha_inicio, fecha_fin):
    """Punto 4: Valida si el servicio está vigente."""
    hoy = datetime.now()

    if not fecha_inicio and not fecha_fin:
        return {
            "ok":           False,
            "alerta":       "sin_datos",
            "detalle":      "❌ No se encontraron fechas de vigencia",
            "fecha_inicio": None,
            "fecha_fin":    None,
        }

    if fecha_fin and fecha_fin < hoy:
        return {
            "ok":           False,
            "alerta":       "vencida",
            "detalle":      (f"⚠️ Servicio VENCIDO el "
                             f"{fecha_fin.strftime('%d/%m/%Y')}"),
            "fecha_inicio": fecha_inicio,
            "fecha_fin":    fecha_fin,
        }

    dias = (fecha_fin - hoy).days if fecha_fin else None

    if dias is not None and dias <= 30:
        return {
            "ok":           True,
            "alerta":       "por_vencer",
            "detalle":      (f"⚠️ Vence en {dias} días "
                             f"({fecha_fin.strftime('%d/%m/%Y')})"),
            "fecha_inicio": fecha_inicio,
            "fecha_fin":    fecha_fin,
        }

    return {
        "ok":           True,
        "alerta":       "vigente",
        "detalle":      (f"✅ Vigente hasta "
                         f"{fecha_fin.strftime('%d/%m/%Y') if fecha_fin else '?'}"),
        "fecha_inicio": fecha_inicio,
        "fecha_fin":    fecha_fin,
    }


def _extraer_tabla_dt_cr(doc):
    """
    Punto 5: Extrae la tabla DT Original.

    Formato real encontrado (Tabla 6 de NEO-0491):
    ['DT Original']
    ['Servicio', 'Tarifa Total*']
    ['Asignación de consultor (1): QA2-Automation QA', 'COL$118,680,000.00']
    ['Asignación de consultor (1): D6-Angle Development', 'COL$131,567,944.00']
    ['Total (1)', 'COL$250,247,944.00']

    También puede estar como párrafo 'DT Original' seguido de tabla.
    """
    filas_resultado = []

    for tabla in doc.tables:
        if not tabla.rows:
            continue

        filas_raw = []
        for fila in tabla.rows:
            celdas = [c.text.strip() for c in fila.cells]
            # Eliminar duplicados por merge
            unicas = []
            prev   = None
            for c in celdas:
                if c != prev:
                    unicas.append(c)
                    prev = c
            if any(unicas):
                filas_raw.append(unicas)

        if not filas_raw:
            continue

        # Verificar si es tabla DT Original
        # Caso 1: primera fila dice 'DT Original'
        es_dt = False
        inicio_datos = 0

        primera = " ".join(filas_raw[0]).lower()
        if "dt original" in primera:
            es_dt        = True
            inicio_datos = 1  # saltar fila DT Original

        # Verificar encabezado servicio/tarifa
        if es_dt and len(filas_raw) > inicio_datos:
            enc = " ".join(filas_raw[inicio_datos]).lower()
            if "servicio" in enc and "tarifa" in enc:
                inicio_datos += 1  # saltar encabezado

        # Caso 2: párrafo anterior dice 'DT Original'
        if not es_dt:
            body   = doc.element.body
            bloques = list(body)
            for idx, bloque in enumerate(bloques):
                tag = bloque.tag.split("}")[-1] if "}" in bloque.tag else ""
                if tag == "tbl" and bloque is tabla._tbl:
                    for j in range(max(0, idx - 4), idx):
                        prev_bloque = bloques[j]
                        prev_texto  = "".join(
                            t.text for t in
                            prev_bloque.iter(f"{{{W}}}t")
                            if t.text
                        ).strip().lower()
                        if "dt original" in prev_texto:
                            es_dt = True
                            # Verificar encabezado
                            if filas_raw and "servicio" in " ".join(
                                    filas_raw[0]).lower():
                                inicio_datos = 1
                            break
                    break

        if not es_dt:
            continue

        # Extraer filas de servicios
        for fila in filas_raw[inicio_datos:]:
            if len(fila) < 2:
                continue

            servicio = fila[0].strip()
            tarifa   = fila[-1].strip()

            # Saltar filas que no son servicios
            if not servicio:
                continue
            if any(x in servicio.lower() for x in [
                "total", "otros costos", "listar", "moneda",
                "servicio", "el total"
            ]):
                continue

            valor_num = limpiar_monto(tarifa)

            filas_resultado.append({
                "servicio":   servicio,
            })

        # Solo procesar la primera tabla DT Original
        if filas_resultado:
            break

    return filas_resultado

def _extraer_tabla_tarifa_servicios(doc):
    """
    Para SOWs SIN CR (primera contratación).
    La tabla de servicios está bajo '1. Tarifa de los Servicios',
    sin el encabezado 'DT Original'.
    """
    body    = doc.element.body
    bloques = list(body)

    # Buscar el párrafo que contiene "Tarifa de los Servicios"
    idx_seccion = None
    for idx, bloque in enumerate(bloques):
        tag = bloque.tag.split("}")[-1] if "}" in bloque.tag else ""
        if tag == "p":
            texto = "".join(
                t.text for t in bloque.iter(f"{{{W}}}t") if t.text
            ).strip().lower()
            if "tarifa de los servicios" in texto:
                idx_seccion = idx
                break

    if idx_seccion is None:
        return []

    # Buscar la primera tabla que venga después del párrafo
    for idx in range(idx_seccion + 1, len(bloques)):
        tag = bloques[idx].tag.split("}")[-1] if "}" in bloques[idx].tag else ""
        if tag != "tbl":
            continue

        # Extraer filas de la tabla
        filas_raw = []
        for tr in bloques[idx].iter(f"{{{W}}}tr"):
            celdas = []
            prev   = None
            for tc in tr.iter(f"{{{W}}}tc"):
                texto_celda = "".join(
                    t.text for t in tc.iter(f"{{{W}}}t") if t.text
                ).strip()
                if texto_celda != prev:
                    celdas.append(texto_celda)
                    prev = texto_celda
            if any(celdas):
                filas_raw.append(celdas)

        if not filas_raw:
            continue

        # Saltar fila de encabezados si existe (Servicio / Tarifa)
        inicio = 0
        if filas_raw and "servicio" in " ".join(filas_raw[0]).lower():
            inicio = 1

        filas_resultado = []
        for fila in filas_raw[inicio:]:
            if not fila:
                continue
            servicio = fila[0].strip()
            if not servicio:
                continue
            if any(x in servicio.lower() for x in [
                "total", "otros costos", "listar", "moneda",
                "servicio", "el total", "tarifa", "descripción"
            ]):
                continue
            filas_resultado.append({"servicio": servicio})

        if filas_resultado:
            return filas_resultado

    return []


def _extraer_tabla_dt_original(doc, numero_sow=None):
    """
    Punto 5: enruta según si el SOW es renovación (tiene CR) o primera vez.
      - Con CR  → busca 'DT Original' en tabla o párrafo previo
      - Sin CR  → busca tabla bajo '1. Tarifa de los Servicios'
    Siempre intenta DT Original primero como fallback para ambos casos.
    """
    es_renovacion = bool(
        re.search(r"-CR\d+", str(numero_sow or ""), re.IGNORECASE)
    )

    if es_renovacion:
        return _extraer_tabla_dt_cr(doc)

    # Sin CR: intentar DT Original por si acaso el doc lo trae
    resultado = _extraer_tabla_dt_cr(doc)
    if resultado:
        return resultado

    # No tiene DT Original → buscar bajo Tarifa de los Servicios
    return _extraer_tabla_tarifa_servicios(doc)


def _validar_clausula_impuestos(texto, doc):
    """
    Punto 6: Busca la cláusula en texto visible Y en SDTs del XML.
    Texto real puede tener campo vacío donde va 'no incluye'.
    """
    # Reconstruir texto completo incluyendo SDTs
    root        = doc.element.body
    partes      = []

    for elemento in root.iter():
        tag = elemento.tag.split("}")[-1] if "}" in elemento.tag else ""
        if tag == "t" and elemento.text and elemento.text.strip():
            partes.append(elemento.text)

    texto_completo = " ".join(partes)

    # Buscar la frase completa con "no incluye"
    patron = re.search(
        r"la\s+tarifa\s+no\s+incluye\s+los\s+impuestos\s+a\s+la\s+venta"
        r"\s+aplicables\s+y\s+cualquier\s+otro\s+impuesto\s+requerido",
        texto_completo, re.IGNORECASE)

    ok = bool(patron)

    if ok:
        detalle = "✅ Cláusula de impuestos correcta"
    else:
        # Verificar si existe la frase pero sin "no incluye"
        patron_parcial = re.search(
            r"la\s+tarifa\s+.*?los\s+impuestos\s+a\s+la\s+venta\s+aplicables",
            texto_completo, re.IGNORECASE | re.DOTALL)
        if patron_parcial:
            detalle = ("❌ La cláusula existe pero falta 'no incluye' — "
                       "verificar el documento")
        else:
            detalle = ("❌ No se encontró la cláusula de impuestos")

    return {"ok": ok, "detalle": detalle}


def analizar_sow_completo(contenido_bytes, numero_sow=None,
                           fecha_inicio=None, fecha_fin=None):
    """
    Análisis completo de un Word de SOW — 6 puntos de validación.
    """
    try:
        from docx import Document

        doc   = Document(io.BytesIO(contenido_bytes))
        texto = extraer_texto_parrafos(doc)

        # Punto 1: Contrato marco
        p1 = _validar_contrato_marco(texto, doc)

        # Punto 2: Tipo de servicio
        p2 = _validar_tipo_servicio(texto)

        # Punto 3: Cargos
        p3 = _extraer_cargos(doc)

        # Punto 4: Vigencia
        if not fecha_inicio or not fecha_fin:
            datos_base = extraer_datos_word(contenido_bytes, numero_sow)
            fi = fecha_inicio or datos_base.get("fecha_inicio")
            ff = fecha_fin    or datos_base.get("fecha_fin")
        else:
            fi = fecha_inicio
            ff = fecha_fin
        p4 = _validar_vigencia(fi, ff)

        # Punto 5: Tabla DT Original
        p5 = _extraer_tabla_dt_original(doc, numero_sow=numero_sow)

        # Punto 6: Cláusula impuestos
        p6 = _validar_clausula_impuestos(texto, doc)

        # Validación global
        global_ok = p1["ok"] and p2["ok"] and p4["ok"] and p6["ok"]

        return {
            "contrato_marco":    p1,
            "tipo_servicio":     p2,
            "cargos":            p3,
            "vigencia":          p4,
            "tabla_dt":          p5,
            "impuestos":         p6,
            "validacion_global": global_ok,
            "error":             None,
        }

    except Exception as e:
        import traceback
        traceback.print_exc()
        return {
            "contrato_marco":    {"ok": False, "detalle": []},
            "tipo_servicio":     {"ok": False, "detalle": ""},
            "cargos":            [],
            "vigencia":          {
                "ok": False, "alerta": "error",
                "detalle": "", "fecha_inicio": None,
                "fecha_fin": None
            },
            "tabla_dt":          [],
            "impuestos":         {"ok": False, "detalle": ""},
            "validacion_global": False,
            "error":             str(e),
        }


def extraer_datos_word(contenido_bytes, numero_sow=None):
    """Extrae fecha inicio, fecha fin y valor COP del Word."""
    try:
        from docx import Document

        doc  = Document(io.BytesIO(contenido_bytes))
        root = doc.element.body

        resultado = {
            "fecha_inicio": None,
            "fecha_fin":    None,
            "valor_cop":    None,
            "fuente":       "",
            "error":        None,
        }

        # Estrategia 1: Content Controls SDT
        texto_inicio = extraer_sdt_por_tag(root, "StartDateSOW")
        texto_fin    = extraer_sdt_por_tag(root, "EndDateSOW")

        if texto_inicio:
            fi = parsear_fecha(texto_inicio)
            if fi:
                resultado["fecha_inicio"] = fi
                resultado["fuente"] += f"SDT StartDateSOW: '{texto_inicio}'. "

        if texto_fin:
            ff = parsear_fecha(texto_fin)
            if ff:
                resultado["fecha_fin"] = ff
                resultado["fuente"] += f"SDT EndDateSOW: '{texto_fin}'. "

        # Estrategia 2: Tablas
        tablas_raw = []
        for tabla in doc.tables:
            filas_tabla = []
            for fila in tabla.rows:
                celdas = []
                for celda in fila.cells:
                    texto_celda = celda.text.strip()
                    sdt_textos  = []
                    for sdt in celda._tc.iter(f"{{{W}}}sdt"):
                        sdtContent = sdt.find(f"{{{W}}}sdtContent")
                        if sdtContent is not None:
                            for t in sdtContent.iter(f"{{{W}}}t"):
                                if t.text and t.text.strip():
                                    sdt_textos.append(t.text.strip())
                    if sdt_textos:
                        texto_celda = " ".join(sdt_textos)
                    celdas.append(texto_celda)
                limpias = []
                prev    = None
                for c in celdas:
                    if c != prev:
                        limpias.append(c)
                        prev = c
                if any(c for c in limpias):
                    filas_tabla.append(limpias)
            if filas_tabla:
                tablas_raw.append(filas_tabla)

        datos_tabla = extraer_datos_de_tablas(tablas_raw, numero_sow)

        if not resultado["fecha_inicio"] and datos_tabla["fecha_inicio"]:
            resultado["fecha_inicio"] = datos_tabla["fecha_inicio"]
        if not resultado["fecha_fin"] and datos_tabla["fecha_fin"]:
            resultado["fecha_fin"] = datos_tabla["fecha_fin"]
        if datos_tabla["valor_cop"]:
            resultado["valor_cop"] = datos_tabla["valor_cop"]

        return resultado

    except Exception as e:
        return {
            "fecha_inicio": None,
            "fecha_fin":    None,
            "valor_cop":    None,
            "fuente":       "",
            "error":        str(e),
        }